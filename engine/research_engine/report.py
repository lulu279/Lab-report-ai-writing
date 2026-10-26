"""Editable Chinese report draft. Fixed numeric tables are constructed from results.json."""
from __future__ import annotations
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def _font(run, size=10.5, bold=False):
    run.font.name = 'Times New Roman'
    run.font.color.rgb = RGBColor(0,0,0)
    run.font.size = Pt(size)
    run.bold = bold
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.insert(0, fonts)
    for name, val in [('ascii','Times New Roman'),('hAnsi','Times New Roman'),('eastAsia','SimSun')]:
        fonts.set(qn('w:'+name),val)


def _p(doc, text, size=10.5):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing=1.5
    p.paragraph_format.first_line_indent=Cm(.74)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _font(p.add_run(text),size)
    return p


def _heading(doc, title, level=1):
    p = doc.add_paragraph()
    p.style = doc.styles[f'Heading {min(level,2)}']
    p.paragraph_format.space_before=Pt(11 if level == 1 else 6)
    p.paragraph_format.space_after=Pt(5)
    _font(p.add_run(title),12,True)
    return p


def _caption(doc, text, before=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(4 if before else 2)
    p.paragraph_format.space_after=Pt(3)
    _font(p.add_run(text),9)


def build_docx(out: Path, results: dict, config, figures_dir: Path,
               discussion: str | None = None) -> None:
    doc = Document()
    section=doc.sections[0]
    section.page_width=Cm(21)
    section.page_height=Cm(29.7)
    section.top_margin=Cm(2.20)
    section.bottom_margin=Cm(2.00)
    section.left_margin=Cm(2.35)
    section.right_margin=Cm(2.15)
    norm=doc.styles['Normal']
    norm.font.name='Times New Roman'
    norm.font.size=Pt(10.5)
    norm.paragraph_format.line_spacing=1.5
    for style in ['Heading 1','Heading 2']:
        doc.styles[style].font.name='Times New Roman'
        doc.styles[style].font.color.rgb=RGBColor(0,0,0)
    title=doc.add_paragraph()
    title.alignment=WD_ALIGN_PARAGRAPH.CENTER
    _font(title.add_run('差示扫描量热法（DSC）数据处理与报告草稿'),16,True)
    sub=doc.add_paragraph()
    sub.alignment=WD_ALIGN_PARAGRAPH.CENTER
    _font(sub.add_run(f"数据集：{config.dataset_label}"),11)
    if config.dataset_type == 'synthetic':
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        _font(p.add_run('【合成测试数据——禁止当作真实实验报告提交】'),11,True)
    else:
        _p(doc,'【自动生成草稿，科学解释、数据来源及格式须经人工核验后使用。】',9)
    _heading(doc, '1 数据来源与范围')
    _p(doc,'本文件由 Research DSC Engine 根据本次输入 CSV 和显式配置自动生成。所有峰温、积分焓与衍生结晶度均来自程序计算，不引用历史报告的数值。')
    _p(doc,'注意：当前原型不解析仪器专用 TXT、不自动划分扫描循环、不检测玻璃化转变温度 Tg。')
    _heading(doc, '2 计算方法')
    _p(doc,'在所指定的温度窗口内，以首尾数据点组的平均热流构建线性基线，并对扣除基线后的热流按实际测量时间进行梯形积分。由于 1 mW·s = 1 mJ，积分面积以 mJ 表示；除以样品质量（mg）后得到 J/g。峰温按基线扣除后的局部极值确定。')
    _p(doc,'结晶度仅在提供完全结晶参考熔融焓，且明确检查了冷结晶影响时计算。若参考焓未核验，计算值为暂估，不能作为已验证的材料性质。')
    _heading(doc, '3 计算结果')
    _caption(doc,'表 1 由本次输入数据计算得到的 DSC 特征量',True)
    table=doc.add_table(rows=1, cols=7)
    table.alignment=WD_TABLE_ALIGNMENT.CENTER
    headers=['样品','阶段','转变','峰温 / °C','ΔH / J·g⁻¹','Xc / %','基线敏感度 / %']
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text=''
        p=cell.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        _font(p.add_run(header),8,True)
    for event in results['events']:
        vals=[event['sample_id'],event['segment'],event['phase'],
              f"{event['peak_temperature_c']:.2f}",f"{event['enthalpy_j_g']:.2f}",
              ('—' if event['crystallinity_pct'] is None else f"{event['crystallinity_pct']:.2f}"),
              f"{event['baseline_sensitivity_pct']:.1f}"]
        for cell,val in zip(table.add_row().cells,vals):
            cell.text=''
            p=cell.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            _font(p.add_run(val),8)
    # Minimal three-line table. No vertical rules.
    tblpr=table._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','bottom','insideH','left','right','insideV']:
        el=OxmlElement('w:'+edge)
        if edge in ('top','bottom'):
            el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'10')
        elif edge=='insideH':
            el.set(qn('w:val'),'nil')
        else:
            el.set(qn('w:val'),'nil')
        borders.append(el)
    tblpr.append(borders)
    header_cell=table.rows[0].cells[0]
    for cell in table.rows[0].cells:
        tcPr=cell._tc.get_or_add_tcPr()
        tcBorders=OxmlElement('w:tcBorders')
        bottom=OxmlElement('w:bottom'); bottom.set(qn('w:val'),'single'); bottom.set(qn('w:sz'),'5')
        tcBorders.append(bottom); tcPr.append(tcBorders)
    if any(e['crystallinity_status']=='computed_reference_unverified' for e in results['events']):
        _p(doc,'注：至少一个样品采用了未独立核验的参考熔融焓，所列结晶度必须补充来源与适用性核验。',9)
    from docx.shared import Inches
    idx=0
    for sample in config.samples:
        fig=figures_dir/f'{sample.sample_id}_overview.png'
        if fig.exists():
            p=doc.add_paragraph()
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(fig),width=Inches(5.05))
            idx+=1
            _caption(doc,f'图 {idx}　{sample.sample_id} 全扫描段热流曲线（原始符号约定）')
    _heading(doc,'4 结果讨论（待人工审核）')
    if discussion:
        _p(doc,'以下是 AI 基于程序计算表撰写的未验证讨论草稿，不代表科学审查通过。')
        for para in discussion.split('\n'):
            if para.strip(): _p(doc,para.strip())
    else:
        _p(doc,'当前未调用 AI 写作。请依据真实样品、测试条件与有效理论补充观察—证据—解释—局限的讨论，不得从历史报告复制数值或结论。')
    _heading(doc,'5 质量状态与限制')
    qa=results['validation']
    _p(doc,f"数值质量门槛：{qa['numerical_status']}；科学解释：{qa['scientific_interpretation_status']}。机器数值检查不等于独立科学审查。")
    _p(doc, f"完整的 {len(qa['issues'])} 条质量提示及数据溯源，请参见同一运行目录的 validation.json、results.json、manifest.json 与 inputs。", 9)
    if any(issue['code']=='REFERENCE_UNVERIFIED' for issue in qa['issues']):
        _p(doc,'待核验：参考熔融焓来源与数值的材料适用性；结晶度仅为基于给定参数的暂估。',9)
    if config.dataset_type == 'synthetic':
        _p(doc,'演示限制：所有曲线均为程序生成的合成数据，绝非仪器测量结果。',9)
    _p(doc,'尚未实现：玻璃化转变的自动 Tg 判断、仪器标定与专有 TXT 解析。基线窗口敏感性不等于仪器测量不确定度。',9)
    foot=section.footer.paragraphs[0]
    foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
    _font(foot.add_run('Research DSC Engine · 未审阅草稿'),9)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
